"""Subir el registro de OT internas que vivía en una planilla.

La planilla trae una línea por ítem y el número de OTI repetido: una misma
orden puede ocupar varias líneas, cada una con su cantidad y su descripción.
Acá se agrupan por número, porque eso es lo que son: una orden con sus ítems.

La fecha venía partida en tres columnas (AÑO, MES, DÍA) y el mes escrito en
palabras, a veces en mayúsculas y a veces no ("Julio", "AGOSTO"). Se arma una
fecha de verdad: así se ordena, se filtra y se compara con las otras dos.
"""

import re
import unicodedata
from datetime import date, datetime

from app.extensions import db
from app.models.oti import ItemOTI, OrdenTrabajoInterna
from app.utils.cantidades import a_cantidad

MESES = (
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
)
# "setiembre" se escribe de las dos formas y las dos son correctas.
ALIAS_MES = {"setiembre": 9}

ESTADOS = OrdenTrabajoInterna.ESTADOS
# Lo que la planilla pone cuando una fórmula no encontró el dato. No es un
# valor: cargarlo dejaría "#N/A" como si fuera una línea de negocio más, y
# apareciendo en el desplegable de filtros.
BASURA = {"#N/A", "#N/D", "#VALUE!", "#REF!", "-", "--"}


def _texto(valor, limite=255) -> str:
    if valor is None:
        return ""
    return re.sub(r"\s+", " ", str(valor)).strip()[:limite]


def _sin_tildes(texto: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn"
    )


def mes_a_numero(valor):
    """El número del mes, venga escrito en palabras o en cifras. None si no se entiende."""
    texto = _sin_tildes(_texto(valor)).lower()
    if not texto:
        return None
    if texto.isdigit():
        n = int(texto)
        return n if 1 <= n <= 12 else None
    if texto in ALIAS_MES:
        return ALIAS_MES[texto]
    for i, nombre in enumerate(MESES, start=1):
        if texto.startswith(nombre[:3]) and nombre.startswith(texto[:3]):
            return i
    return None


def _a_fecha(valor):
    """Una fecha desde una celda, que puede venir como fecha o como texto."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    if not texto:
        return None
    for formato in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def fecha_de_la_fila(anio, mes, dia, *respaldos):
    """La fecha de emisión, armada desde las tres columnas de la planilla.

    Devuelve (fecha, exacta). Cuando el mes o el día vienen en blanco —la
    OTI 350 es así— se usa el primer respaldo que haya, y si no hay
    ninguno, el 1 de enero del año.

    Es una fecha inventada y por eso se devuelve marcada, para informarla.
    Pero perder la orden es peor: su número queda sin dueño en una
    numeración que no tiene saltos, y el trabajo que registraba desaparece
    del historial. Con la fecha a la vista y editable, se corrige mirándola.
    """
    a = None
    try:
        a = int(float(_texto(anio)))
        d = int(float(_texto(dia)))
        m = mes_a_numero(mes)
        if m:
            return date(a, m, d), True
    except (TypeError, ValueError):
        pass
    for respaldo in respaldos:
        if respaldo is not None:
            return respaldo, False
    return (date(a, 1, 1), False) if a else (None, False)


def _estado(valor) -> str:
    """El estado de la planilla, o "SIN ESTADO" si la celda viene vacía.

    382 de las 909 líneas no traen estado, todas entre 2010 y 2022: esa
    columna se empezó a llenar después. Darlas por abiertas llenaría la
    pantalla de pendientes con 240 órdenes en vez de 26, y darlas por
    cerradas sería afirmar que se hicieron. No se sabe, y eso es lo que
    dice: quedan fuera de la vista por defecto y se pueden filtrar aparte.
    """
    texto = _sin_tildes(_texto(valor)).upper()
    return texto if texto in ESTADOS else "SIN ESTADO"


def _limpio(valor, limite=120) -> str:
    """Un texto de la planilla, o vacío si lo que trae no es un dato."""
    texto = _texto(valor, limite)
    return "" if texto.upper() in BASURA else texto


def _columnas(encabezados) -> dict:
    """Dónde está cada columna, buscándola por su nombre.

    Por nombre y no por posición: la planilla tiene una primera columna vacía y
    el orden puede cambiar entre versiones del archivo.
    """
    def buscar(*claves):
        for i, e in enumerate(encabezados):
            nombre = _sin_tildes(_texto(e)).upper()
            if any(c in nombre for c in claves):
                return i
        return None

    return {
        "numero": buscar("OTI", "N OTI"),
        "anio": buscar("ANO", "AÑO"),
        "mes": buscar("MES"),
        "dia": buscar("DIA"),
        "inicio": buscar("FECHA DE INICIO"),
        "estado": buscar("ESTADO"),
        "termino": buscar("TERMINO"),
        "customer": buscar("CUSTOMER"),
        "sucursal": buscar("SUCURSAL"),
        "cantidad": buscar("QTY", "CANTIDAD"),
        "descripcion": buscar("VENTA", "DESCRIPCION", "DETALLE"),
        "ot_qms": buscar("QMS"),
        "cliente": buscar("CLIENTE"),
        "linea": buscar("LINEA"),
    }


def leer_ot_internas(file_storage) -> tuple:
    """Lee la planilla: devuelve (órdenes, problemas).

    Cada orden es un dict con su cabecera y su lista de ítems. 'problemas'
    junta lo que no se pudo leer, para decirlo al final en vez de descartarlo
    en silencio.
    """
    from openpyxl import load_workbook

    libro = load_workbook(file_storage.stream, data_only=True, read_only=True)
    hoja = libro.active
    filas = hoja.iter_rows(values_only=True)
    encabezados = list(next(filas, []))
    col = _columnas(encabezados)
    if col["numero"] is None or col["descripcion"] is None:
        raise ValueError(
            "No encontré las columnas de la planilla. Hacen falta al menos "
            "“Nº OTI” y la del detalle (“VENTA”, “Descripción”)."
        )

    at = lambda fila, clave: (
        fila[col[clave]] if col[clave] is not None and col[clave] < len(fila) else None
    )
    ordenes = {}
    problemas = {"sin_numero": 0, "sin_fecha": [], "sin_detalle": [], "fecha_a_medias": []}
    for fila in filas:
        if not any(c not in (None, "") for c in fila):
            continue
        crudo = _texto(at(fila, "numero"))
        if not crudo or not crudo.replace(".0", "").isdigit():
            problemas["sin_numero"] += 1
            continue
        numero = int(float(crudo))
        inicio = _a_fecha(at(fila, "inicio"))
        termino = _a_fecha(at(fila, "termino"))
        fecha, exacta = fecha_de_la_fila(
            at(fila, "anio"), at(fila, "mes"), at(fila, "dia"), inicio, termino
        )
        if fecha is None:
            problemas["sin_fecha"].append(numero)
            continue
        if not exacta and numero not in problemas["fecha_a_medias"]:
            problemas["fecha_a_medias"].append(numero)

        if numero not in ordenes:
            ordenes[numero] = {
                "numero": numero, "fecha": fecha, "fecha_inicio": inicio,
                "fecha_termino": termino,
                "estado": _estado(at(fila, "estado")),
                "customer": _limpio(at(fila, "customer")),
                "cliente": _limpio(at(fila, "cliente")),
                "sucursal": _limpio(at(fila, "sucursal"), 40),
                "linea_negocio": _limpio(at(fila, "linea"), 40),
                "ot_qms": _limpio(at(fila, "ot_qms"), 30),
                "items": [],
            }
        descripcion = _texto(at(fila, "descripcion"), 500)
        if not descripcion:
            problemas["sin_detalle"].append(numero)
            continue
        ordenes[numero]["items"].append(
            {"cantidad": a_cantidad(at(fila, "cantidad")), "descripcion": descripcion}
        )

    # Una orden sin un solo ítem legible no es una orden: sería un número y una
    # fecha, sin nada que decir qué se mandó a hacer.
    completas = [o for o in ordenes.values() if o["items"]]
    return sorted(completas, key=lambda o: o["numero"]), problemas


def importar_ot_internas(file_storage, empresa_id: int, usuario_id=None) -> dict:
    """Carga la planilla. Las órdenes que ya están no se tocan.

    Se sube una vez, para traer el historial. Volver a subirla no duplica ni
    pisa lo que se haya corregido a mano después: una orden que ya existe se
    deja como está y se informa.
    """
    ordenes, problemas = leer_ot_internas(file_storage)

    existentes = {
        o.numero
        for o in OrdenTrabajoInterna.query.filter_by(empresa_id=empresa_id).all()
    }
    creadas, items = 0, 0
    for datos in ordenes:
        if datos["numero"] in existentes:
            continue
        lineas = datos.pop("items")
        orden = OrdenTrabajoInterna(empresa_id=empresa_id, creado_por_id=usuario_id, **datos)
        for linea in lineas:
            orden.items.append(ItemOTI(**linea))
        db.session.add(orden)
        creadas += 1
        items += len(lineas)

    db.session.commit()
    return {
        "leidas": len(ordenes),
        "creadas": creadas,
        "ya_estaban": len(ordenes) - creadas,
        "items": items,
        **problemas,
    }
