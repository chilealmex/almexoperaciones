"""Subir el conteo físico de a muchos, en vez de fila por fila.

Antes el conteo sólo se podía escribir artículo por artículo en la pantalla de
Stock y conteo. Con tres mil artículos eso es un día de trabajo, y bodega ya
cuenta con una planilla en la mano.

Lo que importa de este importador es lo que NO hace: toca cantidad_fisica y
nada más. Las cantidades de QMS y Defontana, el costo, la unidad y el nombre
son de los otros importadores, y pisarlos acá rompería el cruce.
"""

import io

from werkzeug.datastructures import FileStorage

from app.models.conteo_inventario import ItemConteoInventario
from app.utils.cantidades import a_cantidad
from app.utils.importar_conteo import importar_conteo_fisico


def _fs(texto, nombre="conteo.csv"):
    return FileStorage(stream=io.BytesIO(texto.encode("utf-8")), filename=nombre)


def _item(db, empresa, codigo, **kw):
    datos = dict(empresa_id=empresa.id, codigo=codigo, nombre="PRODUCTO " + codigo,
                 cantidad_qms=10, cantidad_defontana=12, unidad_qms="UN",
                 costo_unitario_qms=1500, en_qms=True, en_defontana=True)
    datos.update(kw)
    item = ItemConteoInventario(**datos)
    db.session.add(item)
    db.session.commit()
    return item


def test_carga_el_conteo_desde_el_archivo(db, empresa):
    _item(db, empresa, "COD-001")
    _item(db, empresa, "COD-002")

    r = importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;7\nCOD-002;3\n"), empresa.id)

    assert r["actualizados"] == 2
    assert [i.cantidad_fisica for i in ItemConteoInventario.query.order_by(ItemConteoInventario.codigo)] == [
        a_cantidad(7), a_cantidad(3)
    ]


def test_no_toca_nada_mas_que_el_conteo(db, empresa):
    """Lo de QMS y Defontana es de los otros importadores."""
    item = _item(db, empresa, "COD-001", nombre="NOMBRE ORIGINAL", cantidad_qms=10,
                 cantidad_defontana=12, unidad_qms="UN", costo_unitario_qms=1500)

    importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;7\n"), empresa.id)

    db.session.refresh(item)
    assert item.cantidad_fisica == a_cantidad(7)
    assert item.cantidad_qms == 10
    assert item.cantidad_defontana == 12
    assert item.nombre == "NOMBRE ORIGINAL"
    assert item.unidad_qms == "UN"
    assert item.costo_unitario_qms == 1500


def test_el_archivo_pisa_el_conteo_anterior(db, empresa):
    """Subir "el conteo actualizado" y que las correcciones no entren sería lo
    contrario de lo que se pide."""
    item = _item(db, empresa, "COD-001", cantidad_fisica=a_cantidad(99))

    r = importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;7\n"), empresa.id)

    db.session.refresh(item)
    assert item.cantidad_fisica == a_cantidad(7)
    assert r["actualizados"] == 1


def test_el_resumen_separa_lo_que_cambio_de_lo_que_ya_estaba(db, empresa):
    """Volver a subir el mismo archivo no debería decir que cambió todo: el
    número sirve para revisar lo que se movió."""
    _item(db, empresa, "COD-001", cantidad_fisica=a_cantidad(7))
    _item(db, empresa, "COD-002", cantidad_fisica=a_cantidad(1))

    r = importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;7\nCOD-002;3\n"), empresa.id)

    assert r["actualizados"] == 1
    assert r["sin_cambio"] == 1
    assert r["total_filas"] == 2


def test_una_fila_sin_cantidad_no_borra_el_conteo(db, empresa):
    """El archivo puede traer todo el maestro con sólo algunas filas contadas."""
    item = _item(db, empresa, "COD-001", cantidad_fisica=a_cantidad(99))

    r = importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;\n"), empresa.id)

    db.session.refresh(item)
    assert item.cantidad_fisica == a_cantidad(99)
    assert r["vacias"] == 1
    assert r["actualizados"] == 0
    # Y se informa como "sin cantidad", no como un número que no se entendió:
    # son cosas distintas y la segunda haría buscar un error que no existe.
    assert r["sin_numero"] == []


def test_queda_registrado_quien_y_cuando(db, empresa, usuario_admin):
    item = _item(db, empresa, "COD-001")

    importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;7\n"), empresa.id, usuario_admin.id)

    db.session.refresh(item)
    assert item.contado_por_id == usuario_admin.id
    assert item.contado_en is not None


def test_un_codigo_que_no_esta_se_informa_y_no_se_crea(db, empresa):
    """El maestro se arma con las planillas de los sistemas, no con el conteo."""
    _item(db, empresa, "COD-001")

    r = importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;7\nNO-EXISTE;5\n"), empresa.id)

    assert r["sin_codigo"] == ["NO-EXISTE"]
    assert ItemConteoInventario.query.count() == 1


def test_acepta_decimales_con_coma(db, empresa):
    """Hay artículos que se cuentan en metros, kilos o litros."""
    item = _item(db, empresa, "COD-001")

    importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;12,5\n"), empresa.id)

    db.session.refresh(item)
    assert item.cantidad_fisica == a_cantidad("12,5")


def test_un_punto_ambiguo_no_se_adivina(db, empresa):
    """"3.125" en Chile es separador de miles tanto como decimal. Elegir mal
    guarda mil veces la cantidad real, sin aviso. Escribiéndolo a mano la
    pantalla pregunta; acá se informa y no se carga."""
    item = _item(db, empresa, "COD-001", cantidad_fisica=a_cantidad(2))

    r = importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;3.125\n"), empresa.id)

    db.session.refresh(item)
    assert item.cantidad_fisica == a_cantidad(2)
    assert r["ambiguos"] == ["COD-001 (3.125)"]
    assert r["actualizados"] == 0


def test_una_cantidad_negativa_no_entra(db, empresa):
    _item(db, empresa, "COD-001")

    r = importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;-3\n"), empresa.id)

    assert r["negativos"] == ["COD-001 (-3)"]
    assert ItemConteoInventario.query.first().cantidad_fisica is None


def test_la_fila_de_totales_no_es_un_articulo(db, empresa):
    """El export del conteo cierra con "Totales (3180)" y la suma de la columna."""
    _item(db, empresa, "COD-001")

    r = importar_conteo_fisico(
        _fs("Código;Stock físico\nCOD-001;7\nTotales (3180);83807\n"), empresa.id
    )

    assert r["sin_codigo"] == []
    assert r["actualizados"] == 1


def test_el_codigo_se_reconoce_escrito_distinto(db, empresa):
    """Bodega escribe el código como le salga: con espacios, en minúsculas."""
    item = _item(db, empresa, "ROP-BCAN-M")

    importar_conteo_fisico(_fs("Código;Stock físico\nrop- bcan-m;7\n"), empresa.id)

    db.session.refresh(item)
    assert item.cantidad_fisica == a_cantidad(7)


def test_un_codigo_retirado_al_unificar_llega_al_articulo_que_quedo(db, empresa):
    """Bodega sigue contando con el código viejo: quien exporta no cambió nada.

    Sin esto, el conteo del código retirado se informaría como "no está en el
    maestro" y se perdería, que es justo lo que la unificación vino a evitar.
    """
    from app.utils.importar_conteo import unificar_grupo

    viejo = _item(db, empresa, "011-CON-OTH-01", en_qms=False, en_defontana=False)
    vigente = _item(db, empresa, "11-CON-OTH-01")
    unificar_grupo([vigente, viejo])
    db.session.commit()

    r = importar_conteo_fisico(_fs("Código;Stock físico\n011-CON-OTH-01;7\n"), empresa.id)

    db.session.refresh(vigente)
    assert r["sin_codigo"] == []
    assert vigente.cantidad_fisica == a_cantidad(7)


def test_si_faltan_las_columnas_lo_dice(db, empresa):
    import pytest

    with pytest.raises(ValueError) as e:
        importar_conteo_fisico(_fs("Otra;Columna\nX;1\n"), empresa.id)
    assert "columnas del conteo" in str(e.value)


def test_el_mismo_codigo_dos_veces_vale_la_ultima(db, empresa):
    item = _item(db, empresa, "COD-001")

    importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;7\nCOD-001;9\n"), empresa.id)

    db.session.refresh(item)
    assert item.cantidad_fisica == a_cantidad(9)


# --- Quién contó y cuándo: mantener el historial ---
#
# Sin estas dos columnas el conteo queda a nombre de quien sube el archivo y con
# la hora de la importación, que es falso en los dos campos: no contó esa persona
# ni fue a esa hora. La exportación de Stock y conteo ya las trae, así que el
# viaje de ida y vuelta —exportar, llenar, subir— tiene que conservarlas.


def test_lee_quien_conto_y_lo_vincula_al_usuario(db, empresa, usuario_admin):
    item = _item(db, empresa, "COD-001")

    importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por\nCOD-001;7;Admin de Prueba\n"), empresa.id
    )

    db.session.refresh(item)
    assert item.contado_por_id == usuario_admin.id
    assert item.contado_por_texto is None
    assert item.quien_conto == "Admin de Prueba"


def test_el_nombre_se_reconoce_escrito_distinto(db, empresa, usuario_admin):
    """Quien llena la planilla escribe como le sale: en mayúsculas, con espacios
    de más, sin tildes."""
    item = _item(db, empresa, "COD-001")

    importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por\nCOD-001;7;  ADMIN  DE  PRUEBA \n"), empresa.id
    )

    db.session.refresh(item)
    assert item.contado_por_id == usuario_admin.id


def test_un_nombre_que_no_es_usuario_se_guarda_igual(db, empresa, usuario_admin):
    """En bodega cuenta gente sin cuenta en el sistema. El nombre escrito es el
    único registro de quién fue: perderlo sería perder el historial, que es
    justo para lo que sirven estas columnas."""
    item = _item(db, empresa, "COD-001")

    r = importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por\nCOD-001;7;Juan Pérez\n"), empresa.id, usuario_admin.id
    )

    db.session.refresh(item)
    assert item.contado_por_id is None
    assert item.contado_por_texto == "Juan Pérez"
    assert item.quien_conto == "Juan Pérez"
    # Se informa igual: casi siempre es un nombre escrito distinto.
    assert r["sin_usuario"] == ["Juan Pérez"]


def test_la_fecha_se_lee_en_hora_de_chile_y_se_guarda_en_utc(db, empresa):
    """La planilla la llena gente en Chile y la exportación sale en hora local,
    pero la base guarda en UTC. Sin convertir, cada conteo quedaría corrido tres
    o cuatro horas, y en el borde del día cambiaría de fecha.

    Las dos filas son del mismo año a propósito: en octubre Chile está en UTC-3
    y en junio en UTC-4. Un desfase fijo acertaría una y fallaría la otra.
    """
    from datetime import datetime, timezone

    _item(db, empresa, "COD-001")
    _item(db, empresa, "COD-002")

    importar_conteo_fisico(_fs(
        "Código;Stock físico;Fecha y hora del conteo\n"
        "COD-001;7;13-10-2026 10:30\n"
        "COD-002;3;20-06-2026 10:30\n"
    ), empresa.id)

    guardadas = {i.codigo: i.contado_en for i in ItemConteoInventario.query}
    def utc(x):
        return x.replace(tzinfo=timezone.utc) if x.tzinfo is None else x.astimezone(timezone.utc)
    assert utc(guardadas["COD-001"]) == datetime(2026, 10, 13, 13, 30, tzinfo=timezone.utc)
    assert utc(guardadas["COD-002"]) == datetime(2026, 6, 20, 14, 30, tzinfo=timezone.utc)


def test_una_celda_de_fecha_de_excel_tambien_se_lee(db, empresa):
    """En .xlsx la fecha llega como fecha, no como texto."""
    from datetime import datetime, timezone
    import openpyxl

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["Código", "Stock físico", "Fecha y hora del conteo"])
    hoja.append(["COD-001", 7, datetime(2026, 10, 13, 10, 30)])
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)

    item = _item(db, empresa, "COD-001")
    importar_conteo_fisico(
        FileStorage(stream=buffer, filename="conteo.xlsx"), empresa.id
    )

    db.session.refresh(item)
    guardada = item.contado_en
    if guardada.tzinfo is None:
        guardada = guardada.replace(tzinfo=timezone.utc)
    assert guardada == datetime(2026, 10, 13, 13, 30, tzinfo=timezone.utc)


def test_sin_esas_columnas_queda_quien_sube_y_la_hora_de_ahora(db, empresa, usuario_admin):
    """El archivo de dos columnas tiene que seguir funcionando igual."""
    item = _item(db, empresa, "COD-001")

    importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;7\n"), empresa.id, usuario_admin.id)

    db.session.refresh(item)
    assert item.contado_por_id == usuario_admin.id
    assert item.contado_por_texto is None
    assert item.contado_en is not None


def test_una_celda_vacia_de_fecha_o_nombre_cae_en_lo_de_siempre(db, empresa, usuario_admin):
    item = _item(db, empresa, "COD-001")

    importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por;Fecha y hora del conteo\nCOD-001;7;;\n"),
        empresa.id, usuario_admin.id,
    )

    db.session.refresh(item)
    assert item.contado_por_id == usuario_admin.id
    assert item.contado_en is not None


def test_una_fecha_que_no_se_entiende_no_bota_la_fila(db, empresa):
    """La cantidad es el dato que importa y ya se leyó bien."""
    item = _item(db, empresa, "COD-001")

    r = importar_conteo_fisico(
        _fs("Código;Stock físico;Fecha y hora del conteo\nCOD-001;7;el martes pasado\n"), empresa.id
    )

    db.session.refresh(item)
    assert item.cantidad_fisica == a_cantidad(7)
    assert item.contado_en is not None
    assert r["fechas_malas"] == ["COD-001 (el martes pasado)"]


def test_contado_por_no_se_confunde_con_la_cantidad(db, empresa):
    """"contado" es palabra de las dos columnas. Si "Contado por" se tomara como
    la cantidad, el conteo entero se cargaría leyendo nombres de personas.

    El encabezado de cantidad acá es "Cantidad" a propósito, sin la palabra
    "físico": con "Stock físico" la columna correcta gana igual por el orden en
    que se buscan, y la confusión no llega a producirse nunca.
    """
    item = _item(db, empresa, "COD-001")

    r = importar_conteo_fisico(
        _fs("Código;Contado por;Cantidad\nCOD-001;Juan Pérez;7\n"), empresa.id
    )

    db.session.refresh(item)
    assert item.cantidad_fisica == a_cantidad(7)
    assert item.contado_por_texto == "Juan Pérez"
    assert r["sin_numero"] == []


def test_volver_a_subir_la_misma_exportacion_no_dice_que_cambio_todo(db, empresa, usuario_admin):
    """La exportación escribe la fecha sin segundos. Sin tolerancia, el segundo
    que se pierde en el viaje haría aparecer como cambio cada una de las filas, y
    el número con que ella revisa lo que se movió dejaría de servir.

    El conteo de partida se guarda con segundos, como queda al contar desde la
    pantalla; el archivo trae el mismo minuto sin ellos, como sale exportado.
    """
    from datetime import datetime, timezone

    _item(db, empresa, "COD-001", cantidad_fisica=a_cantidad(7),
          contado_por_id=usuario_admin.id,
          contado_en=datetime(2026, 10, 13, 13, 30, 45, tzinfo=timezone.utc))
    archivo = ("Código;Stock físico;Contado por;Fecha y hora del conteo\n"
               "COD-001;7;Admin de Prueba;13-10-2026 10:30\n")

    r = importar_conteo_fisico(_fs(archivo), empresa.id, usuario_admin.id)

    assert r["actualizados"] == 0
    assert r["sin_cambio"] == 1


def test_el_nombre_libre_sale_en_el_excel_de_stock(db, empresa, usuario_admin, client):
    """La exportación es el archivo que se vuelve a llenar y subir: si ahí se
    pierde el nombre, el viaje de ida y vuelta lo borra."""
    import openpyxl
    from tests.conftest import login

    _item(db, empresa, "COD-001")
    importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por\nCOD-001;7;Juan Pérez\n"), empresa.id, usuario_admin.id
    )
    login(client, usuario_admin.email)

    datos = client.get("/inventario/stock.xlsx?filtro=todos").data
    hoja = openpyxl.load_workbook(io.BytesIO(datos)).active
    textos = [str(c.value) for fila in hoja.iter_rows() for c in fila]

    assert "Juan Pérez" in textos


def test_corregir_quien_conto_sin_cambiar_la_cantidad_si_entra(db, empresa, usuario_admin):
    """Si sólo estaba mal el nombre, volver a subir el archivo tiene que arreglarlo."""
    item = _item(db, empresa, "COD-001")
    importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por\nCOD-001;7;Juan Pérez\n"), empresa.id, usuario_admin.id
    )

    r = importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por\nCOD-001;7;Pedro Soto\n"), empresa.id, usuario_admin.id
    )

    db.session.refresh(item)
    assert item.contado_por_texto == "Pedro Soto"
    assert r["actualizados"] == 1
    # No es un conteo pisado: la cantidad es la misma, sólo se corrigió el nombre.
    assert r["pisados"] == 0


def test_el_nombre_libre_sobrevive_a_finalizar_la_toma(db, empresa, usuario_admin, client):
    """Si no se archivara, finalizar borraría el nombre justo cuando pasa a ser
    historial, que es para lo que sirve."""
    from app.models.conteo_inventario import TomaInventarioDetalle
    from tests.conftest import login

    _item(db, empresa, "COD-001")
    importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por\nCOD-001;7;Juan Pérez\n"), empresa.id, usuario_admin.id
    )
    login(client, usuario_admin.email)

    client.post("/inventario/toma/cerrar", follow_redirects=True)

    detalle = TomaInventarioDetalle.query.one()
    assert detalle.contado_por_texto == "Juan Pérez"
    assert detalle.quien_conto == "Juan Pérez"
    # Y el cruce vivo quedó limpio, sin el nombre colgando de un conteo borrado.
    assert ItemConteoInventario.query.one().contado_por_texto is None


def test_continuar_una_toma_devuelve_el_nombre_libre(db, empresa, usuario_admin, client):
    from app.models.conteo_inventario import TomaInventario
    from tests.conftest import login

    _item(db, empresa, "COD-001")
    importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por\nCOD-001;7;Juan Pérez\n"), empresa.id, usuario_admin.id
    )
    login(client, usuario_admin.email)
    client.post("/inventario/toma/cerrar", follow_redirects=True)

    toma = TomaInventario.query.one()
    client.post(f"/inventario/historial/{toma.id}/reabrir", follow_redirects=True)

    assert ItemConteoInventario.query.one().contado_por_texto == "Juan Pérez"


def test_el_nombre_libre_se_ve_en_stock_y_conteo(db, empresa, usuario_admin, client):
    from tests.conftest import login

    _item(db, empresa, "COD-001")
    importar_conteo_fisico(
        _fs("Código;Stock físico;Contado por\nCOD-001;7;Juan Pérez\n"), empresa.id, usuario_admin.id
    )
    login(client, usuario_admin.email)

    html = client.get("/inventario/stock?filtro=todos").get_data(as_text=True)

    assert "Juan Pérez" in html


# --- Finalizar la toma anterior antes de subir la captura ---
#
# La captura de bodega es de una toma nueva. Si la anterior no se finalizó, el
# archivo pisa esas cantidades y del conteo anterior no queda rastro: ni en la
# pantalla ni en el historial. No se bloquea la subida —reemplazar la planilla
# de la misma toma, cuando bodega manda una corregida, es legítimo y frecuente—,
# pero tiene que avisarse antes y contarse después.


def test_informa_cuantos_venian_contados_de_la_toma_en_curso(db, empresa):
    """Es lo único que queda de la cantidad anterior una vez pisada."""
    _item(db, empresa, "COD-001", cantidad_fisica=a_cantidad(99))  # venía contado, cambia
    _item(db, empresa, "COD-002", cantidad_fisica=a_cantidad(3))   # venía contado, igual
    _item(db, empresa, "COD-003")                                  # sin contar

    r = importar_conteo_fisico(
        _fs("Código;Stock físico\nCOD-001;7\nCOD-002;3\nCOD-003;5\n"), empresa.id
    )

    assert r["pisados"] == 1
    assert r["actualizados"] == 2
    assert r["sin_cambio"] == 1


def test_con_la_toma_anterior_finalizada_no_se_pisa_nada(db, empresa, usuario_admin, client):
    """El orden que corresponde, de punta a punta: finalizar y después subir.

    Finalizar deja el conteo en blanco, así que la captura entra limpia y no
    pisa nada; y la cantidad anterior queda consultable en el historial, que es
    justamente para lo que sirve finalizar antes.
    """
    from app.models.conteo_inventario import TomaInventarioDetalle
    from tests.conftest import login

    _item(db, empresa, "COD-001", cantidad_fisica=a_cantidad(99))
    login(client, usuario_admin.email)
    client.post("/inventario/toma/cerrar", follow_redirects=True)

    r = importar_conteo_fisico(_fs("Código;Stock físico\nCOD-001;7\n"), empresa.id)

    assert r["pisados"] == 0
    assert r["actualizados"] == 1
    assert TomaInventarioDetalle.query.one().cantidad_fisica == a_cantidad(99)


def test_el_resumen_dice_cuantos_se_pisaron(db, empresa, usuario_admin, client):
    """Después de subir, es el único rastro que queda de la cantidad anterior."""
    from tests.conftest import login

    _item(db, empresa, "COD-001", cantidad_fisica=a_cantidad(99))
    login(client, usuario_admin.email)

    html = client.post(
        "/inventario/conteo/importar/conteo",
        data={"cnt-archivo": _fs("Código;Stock físico\nCOD-001;7\n")},
        content_type="multipart/form-data", follow_redirects=True,
    ).get_data(as_text=True)

    assert "1 ya venían contados en la toma en curso" in html


def test_la_pantalla_avisa_que_hay_una_toma_sin_finalizar(db, empresa, usuario_admin, client):
    """Sin el aviso la mezcla pasa inadvertida: la pantalla no muestra de qué
    toma es cada cantidad."""
    from tests.conftest import login

    _item(db, empresa, "COD-001", cantidad_fisica=a_cantidad(99))
    login(client, usuario_admin.email)

    html = client.get("/inventario/conteo/importar").get_data(as_text=True)

    assert "ya contado en la toma en curso" in html
    assert "Finalizar toma" in html


def test_el_aviso_es_de_la_propia_empresa(db, empresa, usuario_admin, client):
    """Una toma a medio hacer de otra empresa no tiene nada que ver acá: avisaría
    de finalizar algo que en esta pantalla no existe, y no se podría."""
    from app.models.empresa import Empresa
    from tests.conftest import login

    otra = Empresa(rut="77.000.000-0", razon_social="Otra SPA")
    db.session.add(otra)
    db.session.commit()
    _item(db, otra, "COD-001", cantidad_fisica=a_cantidad(99))
    _item(db, empresa, "COD-002")
    login(client, usuario_admin.email)

    html = client.get("/inventario/conteo/importar").get_data(as_text=True)

    assert "en la toma en curso" not in html


def test_sin_nada_contado_no_avisa(db, empresa, usuario_admin, client):
    """El aviso tiene que significar algo: puesto siempre, se deja de leer."""
    from tests.conftest import login

    _item(db, empresa, "COD-001")
    login(client, usuario_admin.email)

    html = client.get("/inventario/conteo/importar").get_data(as_text=True)

    assert "en la toma en curso" not in html
    # Pero la tarjeta sigue ahí
    assert "Conteo físico" in html
