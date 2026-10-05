"""Unificar códigos: el mismo artículo con otro código en cada sistema.

El mismo artículo debería existir en QMS y en Defontana. Cuando el código se
creó distinto en cada uno, el cruce por código deja dos filas sueltas que nunca
se encuentran: una dice "Falta en Defontana" y la otra "Falta en QMS".

Lo que se prueba acá es sobre todo lo que NO debe pasar: que se proponga unir
dos artículos distintos. Unirlos suma sus existencias y falsea el inventario,
y eso es peor que no tener la pantalla.
"""

import io

import pytest
from werkzeug.datastructures import FileStorage

from app.models.conteo_inventario import ItemConteoInventario
from app.models.equivalencia_codigo import EquivalenciaCodigo, traducciones
from app.utils.equivalencias_codigos import (
    UMBRAL,
    esqueleto,
    esqueleto_sin_ceros,
    numeros_significativos,
    proponer,
    puntaje,
)
from tests.conftest import login


class _Articulo:
    """Lo mínimo que mira el motor: un id, un código y un nombre."""

    _n = 0

    def __init__(self, codigo, nombre=""):
        _Articulo._n += 1
        self.id = _Articulo._n
        self.codigo = codigo
        self.nombre = nombre


# --- Normalizar el código ---


@pytest.mark.parametrize("codigo, esperado", [
    ("GOL.PRE-5_8", "GOLPRE58"),
    ("gol pre 5 8", "GOLPRE58"),
    ("CÓDIGO-1", "CODIGO1"),
    ("", ""),
])
def test_el_esqueleto_deja_solo_letras_y_numeros(codigo, esperado):
    assert esqueleto(codigo) == esperado


@pytest.mark.parametrize("codigo, esperado", [
    ("00-FSR-SCW-05", "0FSRSCW5"),
    ("0136-1005", "136-1005".replace("-", "")),
    ("FSR-SCW-5", "FSRSCW5"),
])
def test_los_ceros_de_relleno_no_son_parte_del_codigo(codigo, esperado):
    assert esqueleto_sin_ceros(codigo) == esperado


def test_un_tramo_de_puros_ceros_es_un_cero_y_no_nada():
    """Si "0" desapareciera, "A-0-B" y "A-B" se leerían como el mismo código."""
    assert esqueleto_sin_ceros("A-0-B") == "A0B"
    assert esqueleto_sin_ceros("A-B") == "AB"


@pytest.mark.parametrize("codigo, esperado", [
    ("PERNO-10", (10,)),
    ("00-FSR-SCW-05", (5,)),      # el 00 es relleno: no cuenta
    ("0136-1005", (136, 1005)),
    ("SOLO-LETRAS", ()),
])
def test_los_numeros_significativos_ignoran_el_relleno(codigo, esperado):
    assert numeros_significativos(codigo) == esperado


# --- Los que SÍ son el mismo artículo ---


@pytest.mark.parametrize("uno, otro, nombre", [
    ("GOL.PRE-5_8", "GOLPRE-58", "GOLILLA PRESION 5/8"),
    ("CABLE-3MM", "CABLE_3_MM", "Cable de acero 3mm"),
    ("11-CON-OTH-01", "11CONOTH1", "RELAY E-MECH 6A"),
])
def test_el_mismo_codigo_escrito_distinto_se_reconoce(uno, otro, nombre):
    assert puntaje(_Articulo(uno, nombre), _Articulo(otro, nombre)) == 1.0


def test_el_cero_de_relleno_no_hace_otro_articulo():
    a = _Articulo("0136-1005", "PER.BUTTOM INOX 1/4X5/8")
    b = _Articulo("136-1005", "PER BUTTOM INOX 1/4 X 5/8")
    assert puntaje(a, b) == 1.0


def test_codigos_parecidos_con_el_mismo_nombre_se_proponen():
    a = _Articulo("00-FSR-SCW-05", "SCREWS")
    b = _Articulo("FSR-SCW-5", "SCREWS")
    assert puntaje(a, b) >= UMBRAL


# --- Los que NO son el mismo artículo: lo importante ---


@pytest.mark.parametrize("uno, otro, nombre_a, nombre_b", [
    ("PERNO-10", "PERNO-12", "PERNO HEX 10MM", "PERNO HEX 12MM"),
    ("BROCA-5.0", "BROCA-5.5", "BROCA HSS METRICA 5.0", "BROCA HSS METRICA 5.5"),
    ("SCW-100", "SCW-1000", "TORNILLO 100", "TORNILLO 1000"),
])
def test_dos_medidas_del_mismo_producto_no_se_proponen(uno, otro, nombre_a, nombre_b):
    """El caso peligroso: los códigos se parecen muchísimo y son artículos
    distintos. El parecido carácter a carácter no lo ve —"10" y "12" difieren
    en un dígito de siete— así que hay que mirar los números aparte."""
    valor = puntaje(_Articulo(uno, nombre_a), _Articulo(otro, nombre_b))
    assert valor < UMBRAL, f"{uno} vs {otro} se propondría con {valor}"


def test_dos_articulos_sin_nada_que_ver_no_se_proponen():
    a = _Articulo("TUERCA-8", "TUERCA 8MM")
    b = _Articulo("AMPOLLETA-8", "AMPOLLETA LED 8W")
    assert puntaje(a, b) < UMBRAL


# --- Emparejar ---


def test_cada_articulo_se_propone_una_sola_vez():
    """Ofrecer dos destinos para el mismo artículo garantiza que uno esté mal."""
    qms = [_Articulo("CABLE-3MM", "Cable 3mm")]
    defo = [_Articulo("CABLE_3_MM", "Cable 3 mm"), _Articulo("CABLE-3-MM", "Cable 3mm")]

    propuestas = proponer(qms, defo)

    assert len(propuestas) == 1


def test_las_mas_seguras_van_primero():
    qms = [_Articulo("AAA-1", "Pieza A"), _Articulo("BBB-2", "Pieza B")]
    defo = [_Articulo("BBB2", "Pieza B"), _Articulo("AAA1X", "Pieza A")]

    propuestas = proponer(qms, defo)

    assert propuestas[0]["puntaje"] >= propuestas[-1]["puntaje"]


def test_sin_candidatos_no_inventa_parejas():
    assert proponer([], [_Articulo("X", "X")]) == []
    assert proponer([_Articulo("X", "X")], []) == []


# --- La pantalla ---


def _item(db, empresa, codigo, nombre, en_qms=True, en_defontana=True, **campos):
    item = ItemConteoInventario(
        empresa_id=empresa.id, codigo=codigo, nombre=nombre,
        en_qms=en_qms, en_defontana=en_defontana, **campos,
    )
    db.session.add(item)
    db.session.commit()
    return item


def test_la_pantalla_propone_los_sueltos(client, db, empresa, usuario_admin):
    _item(db, empresa, "GOL.PRE-5_8", "GOLILLA PRESION 5/8", en_defontana=False, cantidad_qms=4)
    _item(db, empresa, "GOLPRE-58", "GOL.PRESION 5-8", en_qms=False, cantidad_defontana=4)
    _item(db, empresa, "EN-LOS-DOS", "Normal")
    login(client, "admin@test.cl")

    cuerpo = client.get("/inventario/equivalencias").get_data(as_text=True)

    assert "GOL.PRE-5_8" in cuerpo and "GOLPRE-58" in cuerpo
    # El que ya cruza no es candidato a nada
    assert "EN-LOS-DOS" not in cuerpo


def test_unir_deja_la_pareja_guardada(client, db, empresa, usuario_admin):
    _item(db, empresa, "GOL.PRE-5_8", "GOLILLA", en_defontana=False)
    _item(db, empresa, "GOLPRE-58", "GOL.PRESION", en_qms=False)
    login(client, "admin@test.cl")

    client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "GOL.PRE-5_8", "codigo_defontana": "GOLPRE-58",
    }, follow_redirects=True)

    equivalencia = EquivalenciaCodigo.query.one()
    assert equivalencia.codigo_qms == "GOL.PRE-5_8"
    assert equivalencia.codigo_defontana == "GOLPRE-58"
    assert equivalencia.creado_por_id == usuario_admin.id


def test_lo_ya_unido_sale_de_las_propuestas(client, db, empresa, usuario_admin):
    _item(db, empresa, "GOL.PRE-5_8", "GOLILLA", en_defontana=False)
    _item(db, empresa, "GOLPRE-58", "GOL.PRESION", en_qms=False)
    login(client, "admin@test.cl")
    client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "GOL.PRE-5_8", "codigo_defontana": "GOLPRE-58",
    }, follow_redirects=True)

    cuerpo = client.get("/inventario/equivalencias").get_data(as_text=True)
    propuestas = cuerpo[cuerpo.index("Parejas propuestas"):cuerpo.index("Ya unidos")]

    assert "GOLPRE-58" not in propuestas
    assert "Ya unidos <span class=\"text-muted\">(1)</span>" in cuerpo


def test_no_se_puede_unir_un_codigo_a_dos_articulos(client, db, empresa, usuario_admin):
    """Si un código de Defontana apuntara a dos de QMS, su stock se duplicaría.

    Por la pantalla ni siquiera se llega a intentarlo: al unir, la fila de
    Defontana se absorbe en la de QMS y deja de estar suelta.
    """
    _item(db, empresa, "QMS-A", "A", en_defontana=False)
    _item(db, empresa, "QMS-B", "B", en_defontana=False)
    _item(db, empresa, "DEFO-X", "X", en_qms=False)
    login(client, "admin@test.cl")
    client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "QMS-A", "codigo_defontana": "DEFO-X"}, follow_redirects=True)

    client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "QMS-B", "codigo_defontana": "DEFO-X"}, follow_redirects=True)

    assert EquivalenciaCodigo.query.count() == 1
    assert EquivalenciaCodigo.query.one().codigo_qms == "QMS-A"


def test_el_modelo_tampoco_deja_unir_dos_veces_el_mismo_codigo(db, empresa):
    """La guarda de fondo, por si alguna vez se llega por otro camino."""
    from app.models.equivalencia_codigo import crear_equivalencia

    crear_equivalencia(empresa.id, "QMS-A", "DEFO-X", None)
    db.session.commit()

    _equivalencia, error = crear_equivalencia(empresa.id, "QMS-B", "DEFO-X", None)

    assert "ya está unido" in error
    assert EquivalenciaCodigo.query.count() == 1


def test_al_unir_se_juntan_las_dos_filas_en_una(client, db, empresa, usuario_admin):
    """Si la fila de Defontana quedara, la próxima importación mandaría su
    stock a la de QMS y esa se quedaría con el saldo viejo para siempre,
    apareciendo como un artículo que ya no está en ningún sistema."""
    _item(db, empresa, "QMS-A", "Perno", en_defontana=False, cantidad_qms=4)
    _item(db, empresa, "DEFO-X", "PERNO HEX", en_qms=False,
          cantidad_defontana=9, unidad_defontana="UN")
    login(client, "admin@test.cl")

    client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "QMS-A", "codigo_defontana": "DEFO-X"}, follow_redirects=True)

    assert ItemConteoInventario.query.count() == 1
    item = ItemConteoInventario.query.one()
    assert item.codigo == "QMS-A"
    assert item.cantidad_qms == 4
    assert item.cantidad_defontana == 9
    assert item.unidad_defontana == "UN"
    assert item.falta_en == ""


def test_al_unir_no_se_pierde_el_conteo_fisico_de_la_fila_que_se_va(
    client, db, empresa, usuario_admin
):
    """Contar es trabajo de bodega: no puede perderse por unir dos códigos."""
    _item(db, empresa, "QMS-A", "Perno", en_defontana=False, cantidad_qms=4)
    _item(db, empresa, "DEFO-X", "PERNO HEX", en_qms=False,
          cantidad_defontana=9, cantidad_fisica=7)
    login(client, "admin@test.cl")

    client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "QMS-A", "codigo_defontana": "DEFO-X"}, follow_redirects=True)

    assert ItemConteoInventario.query.one().cantidad_fisica == 7


def test_un_codigo_que_no_existe_se_rechaza(client, db, empresa, usuario_admin):
    """Los códigos se escriben a mano: un dedazo no puede quedar guardado como
    una equivalencia hacia un artículo que no existe, sin unir nada y sin que
    se note nunca."""
    _item(db, empresa, "QMS-A", "A", en_defontana=False)
    _item(db, empresa, "DEFO-X", "X", en_qms=False)
    login(client, "admin@test.cl")

    cuerpo = client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "QMS-A", "codigo_defontana": "NO-EXISTE"},
        follow_redirects=True).get_data(as_text=True)

    assert "No existe ningún artículo con el código" in cuerpo
    assert EquivalenciaCodigo.query.count() == 0


def test_unir_dos_codigos_que_ya_cruzan_solos_se_rechaza(db, empresa):
    """Con distinto espaciado y mayúsculas ya cruzan desde la importación:
    guardar una equivalencia ahí sería ruido que después confunde."""
    from app.models.equivalencia_codigo import crear_equivalencia

    _equivalencia, error = crear_equivalencia(empresa.id, "MISMO-1", " mismo-1 ", None)

    assert "ya cruzan solos" in error
    assert EquivalenciaCodigo.query.count() == 0


def test_separar_deshace_la_union(client, db, empresa, usuario_admin):
    _item(db, empresa, "QMS-A", "A", en_defontana=False)
    _item(db, empresa, "DEFO-X", "X", en_qms=False)
    login(client, "admin@test.cl")
    client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "QMS-A", "codigo_defontana": "DEFO-X"}, follow_redirects=True)
    equivalencia = EquivalenciaCodigo.query.one()

    client.post(f"/inventario/equivalencias/{equivalencia.id}/deshacer",
                data={}, follow_redirects=True)

    assert EquivalenciaCodigo.query.count() == 0


def test_la_pantalla_la_ve_quien_puede_editar(client, db, empresa, usuario_bodega):
    login(client, "bodega@test.cl")
    assert client.get("/inventario/equivalencias").status_code == 403


def test_el_enlace_esta_en_el_menu(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")
    assert "/inventario/equivalencias" in client.get("/inventario/").get_data(as_text=True)


# --- Lo que hace que valga la pena: el importador la respeta ---


CSV_DEFONTANA = (
    "CodArticulo;Descripción;CodBodega;Nombre Bodega;Saldo Stock\r\n"
    "GOLPRE-58;GOL.PRESION 5-8;01;Casa Matriz;7\r\n"
)


def test_la_union_se_respeta_en_la_siguiente_importacion(
    client, db, empresa, usuario_admin
):
    """Lo que la hace una tabla de traducción y no una fusión de una sola vez.

    Si sólo se fusionaran las filas, la próxima carga volvería a crear los dos
    artículos separados y habría que rehacer el trabajo todos los meses.
    """
    from app.utils.importar_conteo import importar_defontana

    _item(db, empresa, "GOL.PRE-5_8", "GOLILLA PRESION 5/8",
          en_defontana=False, cantidad_qms=4)
    _item(db, empresa, "GOLPRE-58", "GOL.PRESION 5-8",
          en_qms=False, cantidad_defontana=2)
    login(client, "admin@test.cl")
    client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "GOL.PRE-5_8", "codigo_defontana": "GOLPRE-58",
    }, follow_redirects=True)

    importar_defontana(
        FileStorage(stream=io.BytesIO(CSV_DEFONTANA.encode("utf-8")), filename="d.csv"),
        empresa.id,
    )

    # Un solo artículo, con el stock de los dos sistemas en la misma fila
    assert ItemConteoInventario.query.count() == 1
    item = ItemConteoInventario.query.one()
    assert item.codigo == "GOL.PRE-5_8"
    assert item.cantidad_qms == 4
    assert item.cantidad_defontana == 7
    assert item.falta_en == ""


def test_sin_union_la_importacion_los_deja_separados(client, db, empresa, usuario_admin):
    """El contraste: es exactamente el problema que se viene a resolver."""
    from app.utils.importar_conteo import importar_defontana

    _item(db, empresa, "GOL.PRE-5_8", "GOLILLA PRESION 5/8",
          en_defontana=False, cantidad_qms=4)

    importar_defontana(
        FileStorage(stream=io.BytesIO(CSV_DEFONTANA.encode("utf-8")), filename="d.csv"),
        empresa.id,
    )

    assert ItemConteoInventario.query.count() == 2


def test_las_traducciones_son_de_cada_empresa(db, empresa):
    assert traducciones(empresa.id) == {}


# --- Que la pantalla alcance a responder ---
#
# Comparar todos contra todos es cuadrático: con tres mil artículos por lado
# son nueve millones de comparaciones, cada una con dos medidas de parecido, y
# la pantalla no alcanzaba a abrirse. Se indexa por trozos de código.


def _lote(cantidad, patron, desde):
    return [
        _Articulo(patron.format(i=i), f"ARTICULO DE PRUEBA NUMERO {i}")
        for i in range(desde, desde + cantidad)
    ]


def test_el_indice_no_pierde_ninguna_pareja_confiable():
    """El riesgo de indexar es perder parejas buenas por no llegar a medirlas.

    La garantía es concreta y es la que sostiene el umbral: todo lo que puntúa
    0,80 o más se encuentra. Más abajo el índice puede no llegar a medir una
    pareja que comparte muy poco código, y por eso el umbral de la pantalla
    está sobre ese piso: no se ofrece una revisión exhaustiva que no ocurre.

    Las parejas de este caso NO calzan exacto —se diferencian en el "00-" de
    relleno— así que pasan por el emparejado difuso, que es justo el que puede
    dejar cosas fuera. Y cada una lleva señuelos que comparten con ella más
    trozos de código que su pareja verdadera: si el índice se quedara con unos
    pocos candidatos, los señuelos taparían la buena.
    """
    qms, defontana = [], []
    for i in range(40):
        qms.append(_Articulo(f"00-FSR-SCW-{i:03d}", f"SCREWS {i}"))
        defontana.append(_Articulo(f"FSR-SCW-{i}", f"SCREWS {i}"))
        for extra in range(9, 15):
            # Comparte todos los trozos del de QMS y algunos más. La letra
            # evita que el señuelo termine siendo el mismo código que otro de
            # la lista una vez quitados los ceros de relleno.
            defontana.append(
                _Articulo(f"00-FSR-SCW-{i:03d}-K{extra}", f"OTRA MEDIDA {i}-{extra}")
            )
    qms += [_Articulo(f"Q{7000000 + i}", f"SUELTO EN QMS {i}") for i in range(40)]
    defontana += [_Articulo(f"D{3000000 + i}", f"APARTE DEFONTANA {i}") for i in range(40)]

    encontradas = {
        (p["qms"].codigo, p["defontana"].codigo) for p in proponer(qms, defontana, tope=10**6)
    }
    confiables = {
        (a.codigo, b.codigo) for a in qms for b in defontana if puntaje(a, b) >= 0.80
    }

    assert confiables, "el caso de prueba no tiene ninguna pareja confiable"
    assert confiables <= encontradas, f"se perdieron: {sorted(confiables - encontradas)[:5]}"


def test_con_miles_de_articulos_la_pantalla_responde():
    """Guarda de verdad: si alguien vuelve a comparar todo contra todo, esto
    tarda minutos en vez de segundos y la prueba lo delata."""
    import time

    qms = _lote(1200, "QM{i:05d}-A", 0)
    defontana = _lote(1200, "DF{i:05d}-B", 50000)

    inicio = time.perf_counter()
    proponer(qms, defontana)
    tardanza = time.perf_counter() - inicio

    assert tardanza < 5, f"tardó {tardanza:.1f}s: la pantalla no alcanza a abrirse"


# --- El Excel: para repartirse el trabajo fuera de la pantalla ---


def _hojas(client):
    from openpyxl import load_workbook

    respuesta = client.get("/inventario/equivalencias.xlsx")
    assert respuesta.status_code == 200
    libro = load_workbook(io.BytesIO(respuesta.data))
    return {h.title: [[c.value for c in f] for f in libro[h.title].iter_rows()] for h in libro}


def test_el_excel_trae_las_cuatro_situaciones(client, db, empresa, usuario_admin):
    _item(db, empresa, "GOL.PRE-5_8", "GOLILLA PRESION", en_defontana=False, cantidad_qms=4)
    _item(db, empresa, "GOLPRE-58", "GOL.PRESION", en_qms=False, cantidad_defontana=4)
    _item(db, empresa, "SOLO-EN-QMS", "Sin pareja posible", en_defontana=False, cantidad_qms=7)
    _item(db, empresa, "ZZZZZ-999", "Otra cosa distinta", en_qms=False, cantidad_defontana=2)
    login(client, "admin@test.cl")

    hojas = _hojas(client)

    assert sorted(hojas) == ["Por unir", "Sin pareja — Defontana", "Sin pareja — QMS", "Ya unidos"]


def test_lo_que_falta_confirmar_va_con_su_puntaje_y_su_motivo(
    client, db, empresa, usuario_admin
):
    _item(db, empresa, "GOL.PRE-5_8", "GOLILLA PRESION", en_defontana=False, cantidad_qms=4)
    _item(db, empresa, "GOLPRE-58", "GOL.PRESION", en_qms=False, cantidad_defontana=9)
    login(client, "admin@test.cl")

    valores = [v for fila in _hojas(client)["Por unir"] for v in fila]

    assert "GOL.PRE-5_8" in valores and "GOLPRE-58" in valores
    # Los dos stocks, para poder decidir sin volver a la pantalla
    assert 4 in valores and 9 in valores
    assert any("Mismo código" in str(v) for v in valores)


def test_lo_ya_unido_va_aparte(client, db, empresa, usuario_admin):
    _item(db, empresa, "QMS-A", "A", en_defontana=False)
    _item(db, empresa, "DEFO-X", "X", en_qms=False)
    login(client, "admin@test.cl")
    client.post("/inventario/equivalencias/unir", data={
        "codigo_qms": "QMS-A", "codigo_defontana": "DEFO-X"}, follow_redirects=True)

    hojas = _hojas(client)

    assert any("QMS-A" in str(v) for fila in hojas["Ya unidos"] for v in fila)
    # Y sale de lo pendiente: ya se resolvió
    assert not any("QMS-A" in str(v) for fila in hojas["Por unir"] for v in fila)


def test_los_que_no_tienen_pareja_propuesta_tambien_van(client, db, empresa, usuario_admin):
    """Son los que hay que buscar a mano: sin ellos el archivo no sirve para
    repartirse el trabajo, que es para lo que se descarga."""
    _item(db, empresa, "SOLO-EN-QMS", "Sin pareja posible", en_defontana=False, cantidad_qms=7)
    _item(db, empresa, "ZZZZZ-999", "Otra cosa distinta", en_qms=False, cantidad_defontana=2)
    login(client, "admin@test.cl")

    hojas = _hojas(client)

    assert any("SOLO-EN-QMS" in str(v) for fila in hojas["Sin pareja — QMS"] for v in fila)
    assert any("ZZZZZ-999" in str(v) for fila in hojas["Sin pareja — Defontana"] for v in fila)


def test_el_que_tiene_pareja_propuesta_no_se_repite_en_los_sin_pareja(
    client, db, empresa, usuario_admin
):
    """Si saliera en las dos hojas, se revisaría dos veces el mismo artículo."""
    _item(db, empresa, "GOL.PRE-5_8", "GOLILLA PRESION", en_defontana=False)
    _item(db, empresa, "GOLPRE-58", "GOL.PRESION", en_qms=False)
    login(client, "admin@test.cl")

    hojas = _hojas(client)

    assert not any("GOL.PRE-5_8" in str(v) for fila in hojas["Sin pareja — QMS"] for v in fila)
    assert not any("GOLPRE-58" in str(v) for fila in hojas["Sin pareja — Defontana"] for v in fila)


def test_el_excel_y_la_pantalla_proponen_lo_mismo(client, db, empresa, usuario_admin):
    """Si se armaran por separado podrían diferir, y se trabajaría sobre dos
    listas que no son la misma."""
    import re

    for n in range(6):
        _item(db, empresa, f"00-PZA-{n:03d}", f"PIEZA {n}", en_defontana=False)
        _item(db, empresa, f"PZA-{n}", f"PIEZA {n}", en_qms=False)
    login(client, "admin@test.cl")

    pantalla = set(re.findall(
        r'name="codigo_qms" value="([^"]+)"',
        client.get("/inventario/equivalencias").get_data(as_text=True),
    ))
    # Las cuatro primeras filas son el encabezado del informe y la última es
    # la de totales; en medio van los datos.
    excel = {
        fila[0] for fila in _hojas(client)["Por unir"][4:]
        if fila[0] and not str(fila[0]).startswith("Totales")
    }

    assert pantalla and excel == pantalla


def test_el_boton_esta_en_la_pantalla(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")
    assert "/inventario/equivalencias.xlsx" in client.get(
        "/inventario/equivalencias").get_data(as_text=True)


# --- Por qué no se puede unir a mano: cada causa, dicha con su nombre ---
#
# Que un código no esté entre los sueltos tiene causas muy distintas, y cada
# una se arregla en otra parte. Un solo mensaje para todas obligaba a adivinar
# cuál era, revisando a mano un maestro de miles de artículos.


def _intentar_unir(client, qms, defontana):
    return client.post(
        "/inventario/equivalencias/unir",
        data={"codigo_qms": qms, "codigo_defontana": defontana},
        follow_redirects=True,
    ).get_data(as_text=True)


def _articulo(db, empresa, codigo, **campos):
    item = ItemConteoInventario(empresa_id=empresa.id, codigo=codigo, **campos)
    db.session.add(item)
    db.session.commit()
    return item


def test_si_ya_cruzan_lo_dice_en_vez_de_mandar_a_revisar_la_escritura(
    client, db, empresa, usuario_admin
):
    """El caso de ella: el mismo código en los dos sistemas. El cruce ya los
    junta, así que no hay nada que unir, pero el mensaje mandaba a revisar
    cómo estaba escrito."""
    _articulo(db, empresa, "A2000ETQ-CATNOINFLAMABL", en_qms=True, en_defontana=True)
    _articulo(db, empresa, "OTRO-SUELTO", en_qms=False, en_defontana=True)
    login(client, "admin@test.cl")

    cuerpo = _intentar_unir(client, "A2000ETQ-CATNOINFLAMABL", "OTRO-SUELTO")

    assert "ya existe en los dos sistemas" in cuerpo
    assert "no hay nada que unir" in cuerpo


def test_si_el_codigo_esta_en_el_otro_sistema_lo_dice(client, db, empresa, usuario_admin):
    """Poner los dos al revés es el error más fácil de cometer en esta pantalla."""
    _articulo(db, empresa, "SOLO-EN-DEFO", en_qms=False, en_defontana=True)
    _articulo(db, empresa, "SOLO-EN-QMS", en_qms=True, en_defontana=False)
    login(client, "admin@test.cl")

    cuerpo = _intentar_unir(client, "SOLO-EN-DEFO", "SOLO-EN-QMS")

    assert "está en Defontana, no en QMS" in cuerpo
    assert "al revés" in cuerpo


def test_si_el_codigo_esta_dado_de_baja_lo_dice(client, db, empresa, usuario_admin):
    _articulo(db, empresa, "YA-NO-VIENE", en_qms=False, en_defontana=False)
    _articulo(db, empresa, "SUELTO-DEFO", en_qms=False, en_defontana=True)
    login(client, "admin@test.cl")

    cuerpo = _intentar_unir(client, "YA-NO-VIENE", "SUELTO-DEFO")

    assert "ya no viene en ninguna de las dos planillas" in cuerpo


def test_si_el_codigo_no_existe_muestra_los_parecidos(client, db, empresa, usuario_admin):
    """Con miles de artículos, "revisa cómo está escrito" es media hora de
    búsqueda; los parecidos lo resuelven en el mismo aviso."""
    _articulo(db, empresa, "GOLPRE-58", en_qms=True, en_defontana=False)
    _articulo(db, empresa, "GOLPRE-59", en_qms=True, en_defontana=False)
    _articulo(db, empresa, "SUELTO-DEFO", en_qms=False, en_defontana=True)
    login(client, "admin@test.cl")

    cuerpo = _intentar_unir(client, "GOLPRE-5X8", "SUELTO-DEFO")

    assert "No existe ningún artículo" in cuerpo
    assert "GOLPRE-58" in cuerpo and "GOLPRE-59" in cuerpo


def test_si_ya_estaba_unido_dice_con_cual(client, db, empresa, usuario_admin):
    from app.models.equivalencia_codigo import crear_equivalencia

    _articulo(db, empresa, "QMS-A", en_qms=True, en_defontana=False)
    _articulo(db, empresa, "DEFO-A", en_qms=False, en_defontana=True)
    _articulo(db, empresa, "DEFO-B", en_qms=False, en_defontana=True)
    login(client, "admin@test.cl")
    crear_equivalencia(empresa.id, "QMS-A", "DEFO-A", usuario_admin.id, 100, "a mano")
    db.session.commit()
    ItemConteoInventario.query.filter_by(codigo="DEFO-A").delete()
    db.session.commit()

    cuerpo = _intentar_unir(client, "QMS-A", "DEFO-A")

    assert "ya está unido con" in cuerpo


def test_unir_dos_sueltos_de_verdad_sigue_funcionando(client, db, empresa, usuario_admin):
    """El camino bueno no puede romperse por mejorar los mensajes de error."""
    _articulo(db, empresa, "GOL.PRE-5_8", en_qms=True, en_defontana=False)
    _articulo(db, empresa, "GOLPRE-58", en_qms=False, en_defontana=True)
    login(client, "admin@test.cl")

    cuerpo = _intentar_unir(client, "GOL.PRE-5_8", "GOLPRE-58")

    assert "quedaron unidos" in cuerpo
