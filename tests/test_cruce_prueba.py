"""Cruce de prueba: comparar dos planillas sin cargarlas al sistema.

El maestro de artículos es uno solo y lo miran todos los submódulos de
Inventario. Para ver cómo quedaría un cruce no hace falta escribirlo.
"""

import io

from app.models.conteo_inventario import ItemConteoInventario
from app.models.cruce_prueba import CruceDePrueba
from app.models.importacion_inventario import ultimas_importaciones
from app.models.usuario import Usuario
from tests.conftest import login

CSV_QMS = """﻿Sucursal;Linea Negocio;Categoria; Valor Total Stock CLP ;Stock;Stock Critico;Descripción;Unidad;Código Único;ubicacion_bodega; Valor Unitario CLP
Casa Matriz;GOMAS;CAT-A;100000;10;0;PRODUCTO UNO;UN;COD-001;RACK A1; 10,000
Casa Matriz;ACEROS;CAT-B;70000;7;0;PRODUCTO DOS;UN;COD-002;RACK B2; 10,000
"""

CSV_DEFONTANA = (
    "CodArticulo;Descripci\xf3n Art\xedculo;CodBodega;Nombre Bodega;Saldo Stock;Unidad;Costo Unitario\r\n"
    '"COD-001";"PRODUCTO UNO";"BODEGACENTRAL";"BODEGA CENTRAL";"8";"UN";"9.500"\r\n'
    '"COD-003";"PRODUCTO TRES";"BODEGACENTRAL";"BODEGA CENTRAL";"4";"UN";"500"\r\n'
)


def _archivo(texto, nombre, codificacion="utf-8"):
    return (io.BytesIO(texto.encode(codificacion)), nombre)


def _comparar(client, qms=None, defontana=None):
    datos = {}
    if qms is not None:
        datos["archivo_qms"] = qms
    if defontana is not None:
        datos["archivo_defontana"] = defontana
    return client.post(
        "/inventario/cruce-datos/prueba", data=datos, content_type="multipart/form-data"
    )


def _las_dos(client):
    return _comparar(
        client,
        qms=_archivo(CSV_QMS, "qms.csv"),
        defontana=_archivo(CSV_DEFONTANA, "defontana.csv", "cp1252"),
    )


def _codigos(texto):
    import re

    cuerpo = texto[texto.index("<tbody"):texto.index("</tbody>")]
    return re.findall(r'<td class="text-nowrap fw-semibold">([^<]+)</td>', cuerpo)


# --- Lo central: que no toque nada ---


def test_comparar_no_escribe_un_solo_articulo_en_el_maestro(client, db, empresa, usuario_admin):
    """Es el punto entero de esta pantalla: mirar sin modificar."""
    login(client, "admin@test.cl")

    _las_dos(client)

    assert ItemConteoInventario.query.count() == 0


def test_comparar_no_toca_lo_que_ya_estaba_cargado(client, db, empresa, usuario_admin):
    """El maestro lo miran Stock y conteo, Ajuste inventario y los demás: si la
    comparación lo tocara, les cambiaría los datos a todos."""
    db.session.add(ItemConteoInventario(
        empresa_id=empresa.id, codigo="COD-001", nombre="LO QUE YA ESTABA",
        cantidad_qms=1, cantidad_defontana=1, cantidad_fisica=99,
        costo_unitario_qms=777, en_qms=True, en_defontana=True,
    ))
    db.session.commit()
    login(client, "admin@test.cl")

    _las_dos(client)

    item = ItemConteoInventario.query.filter_by(codigo="COD-001").one()
    assert item.nombre == "LO QUE YA ESTABA"
    assert item.cantidad_qms == 1          # la planilla dice 10
    assert item.cantidad_defontana == 1    # la planilla dice 8
    assert item.cantidad_fisica == 99      # el conteo de bodega, intacto
    assert item.costo_unitario_qms == 777
    assert ItemConteoInventario.query.count() == 1   # no creó COD-002 ni COD-003


def test_comparar_no_deja_rastro_de_importacion(client, db, empresa, usuario_admin):
    """La línea de "última importación" dice cuándo entraron datos al sistema.
    Una comparación no los hace entrar, así que no puede moverla."""
    login(client, "admin@test.cl")

    _las_dos(client)

    assert not any(ultimas_importaciones(empresa.id).values())


# --- Y que, sin tocar nada, diga lo mismo que diría importar ---


def test_la_prueba_da_el_mismo_cruce_que_importar_de_verdad(client, db, empresa, usuario_admin):
    """Una comparación que no coincide con lo que pasaría al cargar no sirve
    para decidir: es justamente para eso que se mira antes."""
    login(client, "admin@test.cl")
    _las_dos(client)
    de_la_prueba = _retrato(client, "/inventario/cruce-datos/prueba?filtro=todos&vacios=si")

    # Ahora sí, cargando de verdad por donde se carga
    client.post(
        "/inventario/conteo/importar/qms",
        data={"qms-archivo": _archivo(CSV_QMS, "qms.csv")},
        content_type="multipart/form-data",
    )
    client.post(
        "/inventario/conteo/importar/defontana",
        data={"def-archivo": _archivo(CSV_DEFONTANA, "defontana.csv", "cp1252")},
        content_type="multipart/form-data",
    )
    del_sistema = _retrato(client, "/inventario/cruce-datos?filtro=todos&vacios=si")

    assert de_la_prueba == del_sistema


def _retrato(client, url):
    """Los códigos y los avisos que muestra una de las dos pantallas."""
    import re

    texto = client.get(url).get_data(as_text=True)
    cuerpo = texto[texto.index("<tbody"):texto.index("</tbody>")]
    filas = []
    for fila in cuerpo.split("</tr>"):
        codigo = re.search(r'<td class="text-nowrap fw-semibold">([^<]+)</td>', fila)
        if codigo:
            filas.append((codigo.group(1), sorted(re.findall(r'badge[^>]*>([^<]+)<', fila))))
    return filas


# --- Las cantidades, que es donde un error no se nota mirando ---


def test_las_cantidades_sobreviven_al_guardado(db, empresa):
    """El texto con que se guardan las cantidades lo escribe str(Decimal), no
    una persona: ahí "7.000" son siete con tres decimales. Leerlo con las
    reglas de acá —donde el punto son los miles— lo volvería siete mil, y la
    pantalla mostraría mil veces el stock que hay sin que nada se vea raro:
    ningún aviso cambia de color por eso.
    """
    from decimal import Decimal

    from app.utils.cruce_prueba import comparar, para_guardar

    lectura = {
        "A-1": {"cantidad": Decimal("7.000"), "nombre": "SIETE", "costo": 100},
        "A-2": {"cantidad": Decimal("12.500"), "nombre": "DOCE Y MEDIO", "costo": 100},
        "A-3": {"cantidad": Decimal("3125.000"), "nombre": "TRES MIL CIENTO VEINTICINCO", "costo": 100},
        "A-4": {"cantidad": Decimal("0.200"), "nombre": "UN QUINTO", "costo": 100},
    }

    items = {i.codigo: i for i in comparar(para_guardar(lectura), {})}

    assert items["A-1"].cantidad_qms == Decimal("7")
    assert items["A-2"].cantidad_qms == Decimal("12.5")
    assert items["A-3"].cantidad_qms == Decimal("3125")
    assert items["A-4"].cantidad_qms == Decimal("0.2")


def test_la_prueba_muestra_las_mismas_cantidades_que_importar(client, db, empresa, usuario_admin):
    """Con decimales, que es donde se nota: la planilla trae 12,5 y el cruce
    tiene que decir 12,5 en los dos caminos."""
    qms = CSV_QMS.replace(";10;0;PRODUCTO UNO", ";12,5;0;PRODUCTO UNO")
    login(client, "admin@test.cl")

    _comparar(
        client,
        qms=_archivo(qms, "qms.csv"),
        defontana=_archivo(CSV_DEFONTANA, "defontana.csv", "cp1252"),
    )
    de_la_prueba = _stock_mostrado(client, "/inventario/cruce-datos/prueba.xlsx")

    client.post("/inventario/conteo/importar/qms",
                data={"qms-archivo": _archivo(qms, "qms.csv")},
                content_type="multipart/form-data")
    client.post("/inventario/conteo/importar/defontana",
                data={"def-archivo": _archivo(CSV_DEFONTANA, "defontana.csv", "cp1252")},
                content_type="multipart/form-data")

    del_sistema = ItemConteoInventario.query.filter_by(codigo="COD-001").one()
    assert de_la_prueba["COD-001"] == float(del_sistema.cantidad_qms) * (10000 - 9500)
    assert float(del_sistema.cantidad_qms) == 12.5


def _stock_mostrado(client, url):
    """El impacto en stock de cada artículo: es (dif. costo) x stock, así que
    una cantidad mal leída se ve ahí aunque la columna de stock no esté."""
    from openpyxl import load_workbook

    hoja = load_workbook(io.BytesIO(client.get(url).data)).active
    filas = list(hoja.iter_rows(values_only=True))
    titulos = next(f for f in filas if f and "Código" in f)
    i_cod, i_imp = titulos.index("Código"), titulos.index("Impacto en stock QMS")
    return {
        f[i_cod]: f[i_imp]
        for f in filas[filas.index(titulos) + 1:]
        if f and f[i_cod] and f[i_imp] is not None
    }


# --- Hacen falta los dos lados ---


def test_con_un_solo_archivo_no_hay_cruce(client, db, empresa, usuario_admin):
    """Con uno solo, todo lo del otro sistema saldría como "falta": eso no es
    un cruce sino un listado mal leído."""
    login(client, "admin@test.cl")

    respuesta = _comparar(client, qms=_archivo(CSV_QMS, "qms.csv"))
    cuerpo = client.get(respuesta.headers["Location"]).get_data(as_text=True)

    assert "Falta el archivo de Defontana" in cuerpo
    assert CruceDePrueba.query.count() == 0


def test_un_archivo_que_no_es_planilla_avisa_en_vez_de_reventar(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")

    respuesta = _comparar(
        client,
        qms=_archivo("esto no tiene las columnas\n", "qms.csv"),
        defontana=_archivo(CSV_DEFONTANA, "defontana.csv", "cp1252"),
    )
    cuerpo = client.get(respuesta.headers["Location"]).get_data(as_text=True)

    assert "No se encontraron las columnas" in cuerpo
    assert CruceDePrueba.query.count() == 0


# --- La pantalla ---


def test_la_pantalla_muestra_el_cruce_de_las_dos_planillas(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")
    _las_dos(client)

    cuerpo = client.get("/inventario/cruce-datos/prueba").get_data(as_text=True)

    assert sorted(_codigos(cuerpo)) == ["COD-001", "COD-002", "COD-003"]
    assert "Falta en Defontana" in cuerpo   # COD-002, sólo en QMS
    assert "Falta en QMS" in cuerpo         # COD-003, sólo en Defontana


def test_la_pantalla_avisa_que_no_se_guardo_nada(client, db, empresa, usuario_admin):
    """Lo primero que hay que poder contestar mirándola es si esto es lo que
    tiene el sistema o lo que dicen dos archivos sueltos."""
    login(client, "admin@test.cl")
    _las_dos(client)

    cuerpo = client.get("/inventario/cruce-datos/prueba").get_data(as_text=True)

    assert "sólo una comparación" in cuerpo
    assert "qms.csv" in cuerpo and "defontana.csv" in cuerpo


def test_los_filtros_funcionan_sobre_la_prueba(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")
    _las_dos(client)

    cuerpo = client.get("/inventario/cruce-datos/prueba?q=COD-003&vacios=si").get_data(as_text=True)

    assert _codigos(cuerpo) == ["COD-003"]


def test_el_excel_de_la_prueba_se_descarga(client, db, empresa, usuario_admin):
    from openpyxl import load_workbook

    login(client, "admin@test.cl")
    _las_dos(client)

    respuesta = client.get("/inventario/cruce-datos/prueba.xlsx")

    assert respuesta.status_code == 200
    hoja = load_workbook(io.BytesIO(respuesta.data)).active
    valores = [c.value for fila in hoja.iter_rows() for c in fila]
    assert "COD-001" in valores
    # Y el título deja dicho que no es lo del sistema
    assert any("no cargado al sistema" in str(v) for v in valores if v)


def test_sin_prueba_hecha_manda_a_subirlas(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")

    respuesta = client.get("/inventario/cruce-datos/prueba")

    assert respuesta.headers["Location"].startswith("/inventario/cruce-datos")


# --- Descartarla ---


def test_descartar_borra_la_comparacion(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")
    _las_dos(client)

    client.post("/inventario/cruce-datos/prueba/terminar", data={})

    assert CruceDePrueba.query.count() == 0


# --- De quién es la prueba ---


def test_la_prueba_de_uno_no_le_cambia_la_pantalla_a_otro(client, db, empresa, usuario_admin, roles):
    """Si fuera de la empresa, mirar unos archivos le cambiaría lo que ve el
    resto, que es justo lo que esta pantalla evita."""
    otro = Usuario(
        empresa_id=empresa.id, nombre_completo="Otra Persona", nombre_usuario="otra",
        email="otra@test.cl", rol_id=roles["admin"].id,
    )
    otro.set_password("password123")
    db.session.add(otro)
    db.session.commit()

    login(client, "admin@test.cl")
    _las_dos(client)
    client.get("/logout")

    login(client, "otra@test.cl")
    respuesta = client.get("/inventario/cruce-datos/prueba")

    assert respuesta.headers["Location"].startswith("/inventario/cruce-datos")


def test_volver_a_comparar_reemplaza_la_anterior(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")
    _las_dos(client)
    _las_dos(client)

    assert CruceDePrueba.query.count() == 1


# --- Permisos ---


def test_comparar_exige_permiso_de_edicion(client, db, empresa, usuario_bodega):
    """El cruce se puede mirar sin permiso de edición; subir archivos no."""
    login(client, "bodega@test.cl")

    assert _las_dos(client).status_code == 403
    assert client.get("/inventario/cruce-datos/prueba").status_code == 403


def test_quien_solo_mira_no_ve_el_formulario(client, db, empresa, usuario_bodega):
    login(client, "bodega@test.cl")

    cuerpo = client.get("/inventario/cruce-datos").get_data(as_text=True)

    assert 'name="archivo_qms"' not in cuerpo


# --- Lo que ya estaba resuelto en Unificar códigos ---


def test_la_prueba_respeta_las_parejas_ya_confirmadas(client, db, empresa, usuario_admin):
    """Sin esto mostraría como "Falta en QMS" pares que alguien ya unió, y
    pediría resolver de nuevo un trabajo ya hecho."""
    from app.models.equivalencia_codigo import crear_equivalencia

    db.session.add_all([
        ItemConteoInventario(empresa_id=empresa.id, codigo="COD-001", en_qms=True, en_defontana=False),
        ItemConteoInventario(empresa_id=empresa.id, codigo="COD-003", en_qms=False, en_defontana=True),
    ])
    db.session.commit()
    crear_equivalencia(empresa.id, "COD-001", "COD-003", usuario_admin.id, 100, "a mano")
    db.session.commit()

    login(client, "admin@test.cl")
    _las_dos(client)
    cuerpo = client.get("/inventario/cruce-datos/prueba?vacios=si").get_data(as_text=True)

    # Las dos planillas traen los dos códigos, pero son el mismo artículo
    assert sorted(_codigos(cuerpo)) == ["COD-001", "COD-002"]


# --- La pantalla del sistema sigue mostrando el sistema ---


def test_el_cruce_del_sistema_no_muestra_la_prueba(client, db, empresa, usuario_admin):
    """Son dos cosas distintas y no pueden confundirse: el cruce del sistema
    tiene que seguir mostrando lo cargado, pase lo que pase con la prueba."""
    db.session.add(ItemConteoInventario(
        empresa_id=empresa.id, codigo="SOLO-EN-EL-SISTEMA",
        cantidad_qms=5, cantidad_defontana=5, en_qms=True, en_defontana=True,
    ))
    db.session.commit()
    login(client, "admin@test.cl")
    _las_dos(client)

    cuerpo = client.get("/inventario/cruce-datos").get_data(as_text=True)

    assert _codigos(cuerpo) == ["SOLO-EN-EL-SISTEMA"]
    assert "COD-002" not in _codigos(cuerpo)
