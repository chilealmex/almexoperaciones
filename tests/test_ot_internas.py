"""OT internas: el registro que vivía en una planilla.

Una orden puede tener varios ítems —hasta 32 en un caso de la planilla real—,
así que son dos tablas: la orden y sus ítems. La pantalla abre en las que
siguen vivas, ordenadas por número, porque 723 órdenes de quince años
esconderían las dos docenas que importan hoy.
"""

import io
from datetime import date

import pytest
from werkzeug.datastructures import FileStorage

from app.models.oti import ItemOTI, OrdenTrabajoInterna, siguiente_numero
from app.utils.importar_oti import fecha_de_la_fila, importar_ot_internas, mes_a_numero
from tests.conftest import login


def _oti(db, empresa, numero, estado="ABIERTA", items=(("1", "TRABAJO"),), **kw):
    datos = dict(empresa_id=empresa.id, numero=numero, fecha=date(2026, 1, 15), estado=estado)
    datos.update(kw)
    orden = OrdenTrabajoInterna(**datos)
    for cantidad, descripcion in items:
        orden.items.append(ItemOTI(cantidad=cantidad, descripcion=descripcion))
    db.session.add(orden)
    db.session.commit()
    return orden


def _planilla(filas, encabezados=None):
    """Una planilla como la real: la primera columna vacía y el mes en palabras."""
    import openpyxl

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(encabezados or ["", "AÑO", "MES", "DIA", "Nº OTI", "FECHA DE INICIO", "ESTADO",
                                "FECHA REAL DE TERMINO", "CUSTOMER", "SUCURSAL", "QTY", "VENTA",
                                "N° OT QMS ", "CLIENTE", "LINEA DE NEGOCIO"])
    for fila in filas:
        hoja.append(fila)
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return FileStorage(stream=buffer, filename="oti.xlsx")


FILA = ["", 2026, "ENERO", 15, 900, None, "ABIERTA", None, "SHAW ALMEX", "PRODUCCION",
        2, "RESERVORIO", "6296", "MINERA LOS PELAMBRES", "PRENSAS"]


# --- La lista ---


def test_abre_en_las_que_siguen_vivas(client, db, empresa, usuario_admin):
    """723 órdenes de quince años esconderían las dos docenas que importan hoy."""
    _oti(db, empresa, 100, "CERRADA")
    _oti(db, empresa, 101, "ABIERTA")
    _oti(db, empresa, 102, "DETENIDA")
    _oti(db, empresa, 103, "NULA")
    login(client, "admin@test.cl")

    html = client.get("/inventario/ot-interna").get_data(as_text=True)

    assert ">101<" in html
    assert ">102<" in html
    assert ">100<" not in html
    assert ">103<" not in html


def test_las_sin_estado_no_cuentan_como_pendientes(client, db, empresa, usuario_admin):
    """Vienen de la planilla vieja, que no registraba esa columna antes de 2023.
    No se sabe si se hicieron: darlas por abiertas llenaría la lista de trabajo
    con 216 órdenes de hace quince años."""
    _oti(db, empresa, 100, "SIN ESTADO")
    _oti(db, empresa, 101, "ABIERTA")
    login(client, "admin@test.cl")

    html = client.get("/inventario/ot-interna").get_data(as_text=True)

    assert ">101<" in html
    assert ">100<" not in html


def test_se_pueden_ver_todas(client, db, empresa, usuario_admin):
    _oti(db, empresa, 100, "CERRADA")
    login(client, "admin@test.cl")

    html = client.get("/inventario/ot-interna?estado=todas").get_data(as_text=True)

    assert ">100<" in html


def test_vienen_ordenadas_por_numero(client, db, empresa, usuario_admin):
    _oti(db, empresa, 101)
    _oti(db, empresa, 822)
    _oti(db, empresa, 450)
    login(client, "admin@test.cl")

    html = client.get("/inventario/ot-interna?orden=numero&dir=asc").get_data(as_text=True)

    assert html.index(">101<") < html.index(">450<") < html.index(">822<")


def test_se_busca_por_lo_que_se_mando_a_hacer(client, db, empresa, usuario_admin):
    """De una orden vieja se recuerda qué se pidió, no su número ni su cliente."""
    _oti(db, empresa, 100, items=(("2", "RESERVORIO DE ENFRIAMIENTO"),))
    _oti(db, empresa, 101, items=(("1", "TABLERO ELECTRICO"),))
    login(client, "admin@test.cl")

    html = client.get("/inventario/ot-interna?estado=todas&q=reservorio").get_data(as_text=True)

    assert ">100<" in html
    assert ">101<" not in html


def test_se_busca_por_numero_y_por_cliente(client, db, empresa, usuario_admin):
    _oti(db, empresa, 100, cliente="MINERA LOS PELAMBRES")
    _oti(db, empresa, 101, cliente="MINERA CENTINELA")
    login(client, "admin@test.cl")

    por_cliente = client.get("/inventario/ot-interna?estado=todas&q=centinela").get_data(as_text=True)
    assert ">101<" in por_cliente and ">100<" not in por_cliente

    por_numero = client.get("/inventario/ot-interna?estado=todas&q=100").get_data(as_text=True)
    assert ">100<" in por_numero and ">101<" not in por_numero


def test_se_filtra_por_linea_y_por_anio(client, db, empresa, usuario_admin):
    _oti(db, empresa, 100, linea_negocio="PRENSAS", fecha=date(2025, 3, 1))
    _oti(db, empresa, 101, linea_negocio="FUSION", fecha=date(2026, 3, 1))
    login(client, "admin@test.cl")

    por_linea = client.get("/inventario/ot-interna?estado=todas&linea=FUSION").get_data(as_text=True)
    assert ">101<" in por_linea and ">100<" not in por_linea

    por_anio = client.get("/inventario/ot-interna?estado=todas&anio=2025").get_data(as_text=True)
    assert ">100<" in por_anio and ">101<" not in por_anio


# --- Crear y modificar ---


def test_el_numero_lo_pone_el_sistema_siguiendo_la_numeracion(client, db, empresa, usuario_admin):
    """La planilla va del 100 al 822 sin un solo salto, y esa continuidad es lo
    que permite nombrar una orden por su número sin ambigüedad."""
    _oti(db, empresa, 822)
    login(client, "admin@test.cl")

    client.post("/inventario/ot-interna/nueva", data={
        "fecha": "2026-10-15", "estado": "ABIERTA",
        "item_cantidad": "3", "item_descripcion": "FABRICAR SOPORTE",
    }, follow_redirects=True)

    nueva = OrdenTrabajoInterna.query.filter_by(numero=823).one()
    assert nueva.items[0].descripcion == "FABRICAR SOPORTE"


def test_una_orden_puede_tener_varios_items(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")

    client.post("/inventario/ot-interna/nueva", data={
        "fecha": "2026-10-15", "estado": "ABIERTA",
        "item_cantidad": ["2", "4"],
        "item_descripcion": ["RESERVORIO", "TABLEROS"],
    }, follow_redirects=True)

    orden = OrdenTrabajoInterna.query.one()
    assert [i.descripcion for i in orden.items] == ["RESERVORIO", "TABLEROS"]


def test_las_filas_en_blanco_no_se_guardan(client, db, empresa, usuario_admin):
    """La pantalla ofrece siempre una fila más de la que hay."""
    login(client, "admin@test.cl")

    client.post("/inventario/ot-interna/nueva", data={
        "fecha": "2026-10-15", "estado": "ABIERTA",
        "item_cantidad": ["2", ""],
        "item_descripcion": ["RESERVORIO", ""],
    }, follow_redirects=True)

    assert ItemOTI.query.count() == 1


def test_sin_ningun_item_no_se_crea(client, db, empresa, usuario_admin):
    """Una orden sin ítems no dice qué se manda a hacer."""
    login(client, "admin@test.cl")

    html = client.post("/inventario/ot-interna/nueva", data={
        "fecha": "2026-10-15", "estado": "ABIERTA",
        "item_cantidad": "", "item_descripcion": "",
    }, follow_redirects=True).get_data(as_text=True)

    assert OrdenTrabajoInterna.query.count() == 0
    assert "al menos un ítem" in html


def test_se_puede_modificar_cualquier_dato(client, db, empresa, usuario_admin):
    orden = _oti(db, empresa, 500, "ABIERTA", cliente="MINERA A")
    login(client, "admin@test.cl")

    client.post(f"/inventario/ot-interna/{orden.id}/editar", data={
        "fecha": "2026-02-20", "estado": "CERRADA", "cliente": "MINERA B",
        "fecha_termino": "2026-02-25",
        "item_cantidad": "9", "item_descripcion": "OTRO TRABAJO",
    }, follow_redirects=True)

    db.session.refresh(orden)
    assert orden.estado == "CERRADA"
    assert orden.cliente == "MINERA B"
    assert orden.fecha == date(2026, 2, 20)
    assert orden.fecha_termino == date(2026, 2, 25)
    assert [i.descripcion for i in orden.items] == ["OTRO TRABAJO"]


def test_modificar_no_cambia_el_numero(client, db, empresa, usuario_admin):
    """Es la identidad de la orden: con el número se la nombra en voz alta."""
    orden = _oti(db, empresa, 500)
    login(client, "admin@test.cl")

    client.post(f"/inventario/ot-interna/{orden.id}/editar", data={
        "fecha": "2026-02-20", "estado": "ABIERTA", "numero": "999",
        "item_cantidad": "1", "item_descripcion": "TRABAJO",
    }, follow_redirects=True)

    db.session.refresh(orden)
    assert orden.numero == 500


def test_modificar_pregunta_antes(client, db, empresa, usuario_admin):
    """Se pisa lo que había, y quien la escribió puede no ser quien la cambia."""
    orden = _oti(db, empresa, 500)
    login(client, "admin@test.cl")

    html = client.get(f"/inventario/ot-interna/{orden.id}/editar").get_data(as_text=True)

    assert "¿Seguro de modificar" in html


def test_crear_no_pregunta(client, db, empresa, usuario_admin):
    """Al crear no hay nada que perder: si sale mal, se modifica."""
    login(client, "admin@test.cl")

    html = client.get("/inventario/ot-interna/nueva").get_data(as_text=True)

    assert "¿Seguro de modificar" not in html


def test_sin_estado_no_se_puede_elegir_a_mano(client, db, empresa, usuario_admin):
    """Lo pone sólo la carga del historial. Elegirlo sería crear a propósito una
    orden sin decir en qué va."""
    login(client, "admin@test.cl")

    html = client.get("/inventario/ot-interna/nueva").get_data(as_text=True)
    opciones = html[html.index('name="estado"'):]

    assert "SIN ESTADO" not in opciones[:opciones.index("</select>")]


def test_una_orden_de_otra_empresa_no_se_puede_abrir(client, db, empresa, usuario_admin):
    from app.models.empresa import Empresa

    otra = Empresa(rut="77.000.000-0", razon_social="Otra SPA")
    db.session.add(otra)
    db.session.commit()
    ajena = _oti(db, otra, 100)
    login(client, "admin@test.cl")

    assert client.get(f"/inventario/ot-interna/{ajena.id}/editar").status_code == 404


# --- La planilla histórica ---


def test_se_carga_la_planilla(client, db, empresa, usuario_admin):
    r = importar_ot_internas(_planilla([FILA]), empresa.id)

    orden = OrdenTrabajoInterna.query.one()
    assert r["creadas"] == 1
    assert orden.numero == 900
    assert orden.fecha == date(2026, 1, 15)
    assert orden.cliente == "MINERA LOS PELAMBRES"
    assert orden.linea_negocio == "PRENSAS"
    assert orden.items[0].descripcion == "RESERVORIO"


def test_las_lineas_del_mismo_numero_son_una_sola_orden(client, db, empresa, usuario_admin):
    """El número se repite en la planilla: es la misma orden con varios ítems."""
    otra = list(FILA)
    otra[11] = "TABLEROS"
    r = importar_ot_internas(_planilla([FILA, otra]), empresa.id)

    assert r["creadas"] == 1
    assert OrdenTrabajoInterna.query.one().items[1].descripcion == "TABLEROS"


def test_la_celda_de_estado_vacia_queda_sin_estado(client, db, empresa, usuario_admin):
    """382 de las 909 líneas no traen estado, todas entre 2010 y 2022."""
    sin = list(FILA)
    sin[6] = None
    importar_ot_internas(_planilla([sin]), empresa.id)

    assert OrdenTrabajoInterna.query.one().estado == "SIN ESTADO"


def test_sin_mes_ni_dia_la_orden_no_se_pierde(client, db, empresa, usuario_admin):
    """La OTI 350 de la planilla real es así. Perderla dejaría su número sin
    dueño en una numeración que no tiene saltos."""
    coja = list(FILA)
    coja[2] = coja[3] = None
    r = importar_ot_internas(_planilla([coja]), empresa.id)

    orden = OrdenTrabajoInterna.query.one()
    assert orden.fecha == date(2026, 1, 1)
    assert r["fecha_a_medias"] == [900]


def test_volver_a_subirla_no_duplica_ni_pisa(client, db, empresa, usuario_admin):
    """Lo que se corrigió a mano después no se puede perder por volver a subir
    el mismo archivo."""
    importar_ot_internas(_planilla([FILA]), empresa.id)
    orden = OrdenTrabajoInterna.query.one()
    orden.cliente = "CORREGIDO A MANO"
    db.session.commit()

    r = importar_ot_internas(_planilla([FILA]), empresa.id)

    assert r["creadas"] == 0
    assert r["ya_estaban"] == 1
    assert OrdenTrabajoInterna.query.one().cliente == "CORREGIDO A MANO"


def test_el_correlativo_sigue_donde_quedo_la_planilla(client, db, empresa, usuario_admin):
    importar_ot_internas(_planilla([FILA]), empresa.id)

    assert siguiente_numero(empresa.id) == 901


def test_el_mes_se_entiende_escrito_de_cualquier_forma(client, db, empresa, usuario_admin):
    """La planilla lo trae como "Julio", "AGOSTO", "Septiembre"..."""
    assert mes_a_numero("Julio") == 7
    assert mes_a_numero("AGOSTO") == 8
    assert mes_a_numero("setiembre") == 9
    assert mes_a_numero("DICIEMBRE") == 12
    assert mes_a_numero("3") == 3
    assert mes_a_numero("cualquier cosa") is None


def test_una_fecha_armada_se_informa_como_tal(client, db, empresa, usuario_admin):
    """La que el sistema inventó tiene que distinguirse de la que venía."""
    completa, exacta = fecha_de_la_fila(2026, "ENERO", 15)
    assert completa == date(2026, 1, 15) and exacta is True

    a_medias, exacta = fecha_de_la_fila(2026, None, None)
    assert a_medias == date(2026, 1, 1) and exacta is False


def test_la_basura_de_las_formulas_no_entra(client, db, empresa, usuario_admin):
    """La planilla trae "#N/A" donde una fórmula no encontró el dato. Cargarlo
    lo dejaría como una línea de negocio más, y saldría en el desplegable."""
    con_basura = list(FILA)
    con_basura[14] = "#N/A"
    importar_ot_internas(_planilla([con_basura]), empresa.id)

    assert OrdenTrabajoInterna.query.one().linea_negocio == ""


def test_una_planilla_que_no_es_la_que_corresponde_lo_dice(client, db, empresa, usuario_admin):
    import pytest

    with pytest.raises(ValueError) as e:
        importar_ot_internas(_planilla([["X", 1]], encabezados=["OTRA", "COSA"]), empresa.id)

    assert "columnas de la planilla" in str(e.value)


# --- El menú ---


def test_aparece_en_el_menu_de_inventario(client, db, empresa, usuario_admin):
    login(client, "admin@test.cl")

    html = client.get("/inventario/stock").get_data(as_text=True)

    assert "OT interna" in html
    assert "/inventario/ot-interna" in html


# --- Quién puede qué ---
#
# Crear una orden es trabajo de todos los días y lo puede hacer quien tenga
# permiso de editar inventario. Cambiar una que ya existe o borrarla, no: lo que
# decía queda reemplazado sin rastro, y quien la escribió puede no ser quien la
# está cambiando.


@pytest.fixture()
def usuario_produccion(db, empresa, roles):
    """Alguien que puede editar inventario pero no es administrador.

    El de bodega no sirve para esto: no tiene permiso de editar, así que no
    distingue "no puede modificar porque no es admin" de "no puede nada".
    """
    from app.models.permiso import RolModuloPermiso
    from app.models.usuario import Rol, Usuario

    rol = Rol(clave="produccion", nombre="Producción")
    db.session.add(rol)
    db.session.commit()
    db.session.add(RolModuloPermiso(rol_id=rol.id, modulo="inventario",
                                    puede_ver=True, puede_editar=True))
    usuario = Usuario(empresa_id=empresa.id, nombre_completo="Produccion de Prueba",
                      nombre_usuario="produccion", email="produccion@test.cl", rol_id=rol.id)
    usuario.set_password("password123")
    db.session.add(usuario)
    db.session.commit()
    return usuario


def test_quien_no_es_admin_si_puede_crear(client, db, empresa, usuario_produccion):
    login(client, "produccion@test.cl")

    client.post("/inventario/ot-interna/nueva", data={
        "fecha": "2026-10-15", "estado": "ABIERTA",
        "item_cantidad": "1", "item_descripcion": "TRABAJO",
    }, follow_redirects=True)

    assert OrdenTrabajoInterna.query.count() == 1


def test_quien_no_es_admin_no_puede_modificar(client, db, empresa, usuario_produccion):
    orden = _oti(db, empresa, 500, cliente="EL ORIGINAL")
    login(client, "produccion@test.cl")

    assert client.get(f"/inventario/ot-interna/{orden.id}/editar").status_code == 403

    client.post(f"/inventario/ot-interna/{orden.id}/editar", data={
        "fecha": "2026-02-20", "estado": "CERRADA", "cliente": "CAMBIADO",
        "item_cantidad": "1", "item_descripcion": "OTRO",
    })

    db.session.refresh(orden)
    assert orden.cliente == "EL ORIGINAL"


def test_quien_no_es_admin_no_puede_eliminar(client, db, empresa, usuario_produccion):
    """Cerrar la puerta de la pantalla no basta: la dirección se puede escribir."""
    orden = _oti(db, empresa, 500)
    login(client, "produccion@test.cl")

    assert client.post(f"/inventario/ot-interna/{orden.id}/eliminar").status_code == 403
    assert OrdenTrabajoInterna.query.count() == 1


def test_a_quien_no_es_admin_no_se_le_ofrecen_los_botones(client, db, empresa, usuario_produccion):
    """Un botón que contesta 403 es peor que no tenerlo: hace perder el trabajo
    escrito antes de descubrir que no se podía."""
    _oti(db, empresa, 500)
    login(client, "produccion@test.cl")

    html = client.get("/inventario/ot-interna").get_data(as_text=True)

    assert "Modificar" not in html
    assert "Eliminar" not in html
    # Pero crear sí, que es lo que sí puede hacer.
    assert "Nueva OT interna" in html


def test_el_admin_puede_eliminar(client, db, empresa, usuario_admin):
    orden = _oti(db, empresa, 500, items=(("2", "UNO"), ("3", "DOS")))
    login(client, "admin@test.cl")

    client.post(f"/inventario/ot-interna/{orden.id}/eliminar", follow_redirects=True)

    assert OrdenTrabajoInterna.query.count() == 0
    # Los ítems se van con ella: dejarlos colgando sería basura invisible.
    assert ItemOTI.query.count() == 0


def test_eliminar_pregunta_antes(client, db, empresa, usuario_admin):
    """No se puede deshacer."""
    orden = _oti(db, empresa, 500)
    login(client, "admin@test.cl")

    html = client.get("/inventario/ot-interna").get_data(as_text=True)

    # Enganchado al envío del formulario, no sólo escrito en la página: el texto
    # puede estar y el botón borrar sin preguntar nada.
    envio = html[html.index('oti/500/eliminar') if 'oti/500/eliminar' in html
                 else html.index('/eliminar'):]
    envio = envio[:envio.index('</form>')]
    assert "onsubmit=" in envio
    assert "¿Eliminar la OT interna N° 500?" in envio
    assert "no se puede deshacer" in envio
    # Y dice qué hacer en el caso que de verdad se da seguido.
    assert "Nula" in envio


def test_no_se_elimina_la_de_otra_empresa(client, db, empresa, usuario_admin):
    from app.models.empresa import Empresa

    otra = Empresa(rut="77.000.000-0", razon_social="Otra SPA")
    db.session.add(otra)
    db.session.commit()
    ajena = _oti(db, otra, 100)
    login(client, "admin@test.cl")

    assert client.post(f"/inventario/ot-interna/{ajena.id}/eliminar").status_code == 404
    assert OrdenTrabajoInterna.query.count() == 1


# --- El Excel ---


def _hoja(datos):
    import openpyxl

    return openpyxl.load_workbook(io.BytesIO(datos)).active


def _textos(hoja):
    return [str(c.value) for fila in hoja.iter_rows() for c in fila]


def test_se_descarga_el_excel(client, db, empresa, usuario_admin):
    _oti(db, empresa, 500, cliente="MINERA A", items=(("2", "RESERVORIO"),))
    login(client, "admin@test.cl")

    respuesta = client.get("/inventario/ot-interna.xlsx")

    assert respuesta.status_code == 200
    textos = _textos(_hoja(respuesta.data))
    assert "500" in textos
    assert "RESERVORIO" in textos
    assert "MINERA A" in textos


def test_el_excel_trae_una_fila_por_item(client, db, empresa, usuario_admin):
    """Es como viene la planilla de la que salió el módulo, y es lo que deja
    sumar cantidades o armar una tabla dinámica."""
    _oti(db, empresa, 500, items=(("2", "UNO"), ("3", "DOS"), ("4", "TRES")))
    login(client, "admin@test.cl")

    hoja = _hoja(client.get("/inventario/ot-interna.xlsx").data)
    con_el_numero = [f for f in hoja.iter_rows(values_only=True) if f and f[0] == 500]

    assert len(con_el_numero) == 3
    assert [f[4] for f in con_el_numero] == ["UNO", "DOS", "TRES"]


def test_el_excel_respeta_los_filtros(client, db, empresa, usuario_admin):
    """El archivo descargado es lo que se está mirando, no otra cosa."""
    _oti(db, empresa, 500, "ABIERTA", items=(("1", "LA VIVA"),))
    _oti(db, empresa, 501, "CERRADA", items=(("1", "LA CERRADA"),))
    login(client, "admin@test.cl")

    por_defecto = _textos(_hoja(client.get("/inventario/ot-interna.xlsx").data))
    assert "LA VIVA" in por_defecto
    assert "LA CERRADA" not in por_defecto

    todas = _textos(_hoja(client.get("/inventario/ot-interna.xlsx?estado=todas").data))
    assert "LA CERRADA" in todas


def test_el_excel_lo_puede_bajar_quien_solo_mira(client, db, empresa, usuario_bodega):
    """Descargar no cambia nada: alcanza con poder ver el módulo."""
    _oti(db, empresa, 500)
    login(client, "bodega@test.cl")

    assert client.get("/inventario/ot-interna.xlsx").status_code == 200


def test_el_boton_de_excel_se_lleva_los_filtros(client, db, empresa, usuario_admin):
    _oti(db, empresa, 500, "CERRADA")
    login(client, "admin@test.cl")

    html = client.get("/inventario/ot-interna?estado=todas").get_data(as_text=True)

    assert "ot-interna.xlsx?estado=todas" in html
